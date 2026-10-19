"""Excel ingestion is deterministic and never invokes an LLM."""
import csv
import io
import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.services.classification import rule_classify

ALIASES = {
    'tx_no': ['交易单号', '交易号', '订单号', 'tx_no'],
    'tx_time': ['交易时间', '付款时间', '时间', '日期', 'tx_time'],
    'merchant': ['交易对方', '商户', '商家', 'merchant'],
    'product': ['商品', '商品名称', '备注', 'product'],
    'amount': ['金额', '金额(元)', '交易金额', '收支金额', 'amount'],
    'direction': ['收/支', '收支', '收支类型', 'direction'],
    'transaction_type': ['交易类型', 'transaction_type'],
    'payment_method': ['支付方式', '付款方式', 'payment_method'],
    'transaction_status': ['当前状态', '交易状态', 'transaction_status'],
    'merchant_tx_no': ['商户单号', 'merchant_tx_no'],
    'note': ['备注', 'note'],
    'category': ['分类', 'category'],
}
DIRECTIONS = {'支出': 'expense', '收入': 'income', '不计收支': 'neutral',
              '不计收入': 'neutral', '不计支出': 'neutral', '/': 'neutral',
              '其他': 'neutral', 'expense': 'expense', 'income': 'income', 'neutral': 'neutral'}


def normalize(value):
    return re.sub(r'\s+', '', str(value or '')).replace('（', '(').replace('）', ')')


def text(value):
    return '' if value is None else str(value).strip()


def parse_amount(value):
    try:
        amount = Decimal(re.sub(r'[,¥￥\s]', '', str(value)))
    except InvalidOperation:
        raise ValueError('金额不是有效数字') from None
    if not amount.is_finite() or amount < 0 or amount > Decimal('9999999999.99'):
        raise ValueError('金额必须是有限的非负数字，且不超过 9999999999.99')
    if amount != amount.quantize(Decimal('0.01')):
        raise ValueError('金额不能多于两位小数')
    return int(amount * 100)


def parse_date(value):
    if isinstance(value, datetime):
        result = value
    else:
        try:
            result = datetime.fromisoformat(text(value).replace('/', '-'))
        except ValueError:
            raise ValueError('时间应为 YYYY-MM-DD HH:MM:SS 或 Excel 日期单元格') from None
    if result.tzinfo:
        raise ValueError('请使用本地无时区时间；本框架账单时区为 Asia/Shanghai')
    return result.isoformat(timespec='seconds')


def parse_tx_no(cell):
    value = cell.value if hasattr(cell, 'value') else cell
    if isinstance(value, (int, float)):
        if value != int(value) or len(str(int(value))) > 15:
            raise ValueError('交易号已被 Excel 当作数值，可能损失精度；请使用原始文本交易号')
        value = str(int(value))
    result = text(value)
    if not result:
        raise ValueError('缺少交易号；当前版本不猜测无交易号账单的去重键')
    if len(result) > 128:
        raise ValueError('交易号过长')
    return result


def parse_rows(rows, sheet, source):
    mapping = None
    valid, errors = [], []
    for number, cells in enumerate(rows, start=1):
        if number > 10030:
            raise ValueError('单个工作表最多支持 10000 行账单')
        values = [cell.value if hasattr(cell, 'value') else cell for cell in cells]
        if not any(value is not None and text(value) for value in values):
            continue
        if mapping is None:
            headers = [normalize(value) for value in values]
            mapping_candidate = {key: next((headers.index(normalize(alias)) for alias in aliases
                                           if normalize(alias) in headers), None)
                                 for key, aliases in ALIASES.items()}
            if all(mapping_candidate[key] is not None for key in ('tx_no', 'tx_time', 'amount', 'direction')):
                mapping = mapping_candidate
                continue
            if number >= 30:
                break
            continue

        def get(key):
            index = mapping[key]
            return cells[index] if index is not None and index < len(cells) else None

        def value(key):
            cell = get(key)
            return cell.value if hasattr(cell, 'value') else cell

        try:
            for key in ALIASES:
                if getattr(get(key), 'data_type', None) == 'f':
                    raise ValueError('关键字段包含公式；请提供原始值')
            direction = DIRECTIONS.get(text(value('direction')))
            if direction is None:
                raise ValueError('无法识别收支类型')
            merchant, product = text(value('merchant')), text(value('product'))
            transaction_type = text(value('transaction_type')) or '未知'
            category = text(value('category')) or '其他'
            extra = {key: text(value(key)) or default for key, default in (
                ('payment_method', '未知'), ('transaction_status', '未知'), ('merchant_tx_no', ''), ('note', '')
            )}
            if any(len(extra[key]) > limit for key, limit in (
                ('payment_method', 64), ('transaction_status', 64), ('merchant_tx_no', 128), ('note', 500)
            )):
                raise ValueError('支付方式、状态、商户单号或备注文本过长')
            if (len(merchant) > 255 or len(product) > 500 or len(transaction_type) > 64
                    or len(category) > 64):
                raise ValueError('商户、商品或分类文本过长')
            row = {
                'tx_no': parse_tx_no(get('tx_no')), 'tx_time': parse_date(value('tx_time')),
                'merchant': merchant, 'product': product, 'amount_cents': parse_amount(value('amount')),
                'direction': direction, 'transaction_type': transaction_type, 'category': category,
                'source': source,
                'source_sheet': sheet, 'source_row': number,
                **extra,
            }
            if text(value('category')):
                row.update(category_source='source', category_confidence=1.0)
            else:
                row.update(rule_classify(row))
            valid.append(row)
        except (ValueError, TypeError, InvalidOperation) as exc:
            errors.append({'sheet': sheet, 'row': number, 'message': str(exc)})
    if mapping is None:
        errors.append({'sheet': sheet, 'row': 0, 'message': '前 30 行未找到交易号、交易时间、金额、收支表头'})
    return valid, errors


def parse_file(content: bytes, filename: str, source: str):
    if filename.lower().endswith('.csv'):
        decoded = None
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ValueError('无法识别 CSV 编码')
        valid, errors = parse_rows(csv.reader(io.StringIO(decoded)), 'CSV', source)
    elif filename.lower().endswith('.xlsx'):
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 64 * 1024**2:
                    raise ValueError('Excel 解压后过大')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except (BadZipFile, KeyError):
            raise ValueError('不是有效的 .xlsx 文件') from None
        valid, errors = [], []
        try:
            if len(workbook.worksheets) > 20:
                raise ValueError('最多支持 20 个工作表')
            for sheet in workbook.worksheets:
                if sheet.max_column and sheet.max_column > 100:
                    raise ValueError('工作表超过 100 列')
                rows, issues = parse_rows(sheet.iter_rows(), sheet.title, source)
                valid.extend(rows)
                errors.extend(issues)
                if len(valid) + len(errors) > 10000:
                    raise ValueError('一次导入最多支持 10000 行')
        finally:
            workbook.close()
    else:
        raise ValueError('只支持 .xlsx 或 .csv')
    return {'rows': valid, 'errors': errors, 'status': 'preview'}


def serialize(data):
    return json.dumps(data, ensure_ascii=False)
